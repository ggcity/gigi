#!/usr/bin/env python3
"""
Website Crawler for the City Website RAG System (local edition).

Crawls a website, extracts text + metadata, embeds chunks locally with
BGE-large, and stores them in a local Chroma collection. No AWS.

Notes:
  * Embeddings are computed in a worker thread (asyncio.to_thread) so the
    CPU-bound encode does not block concurrent downloads on the event loop.
  * Chunks of a page are embedded in one batched call for CPU throughput.
  * Chroma accepts custom IDs, so re-crawling upserts by a deterministic
    per-chunk ID instead of wiping (unless --recreate is passed).
  * The retention-analysis summary is computed in Python over Chroma metadata,
    since Chroma has no server-side aggregations.
"""

import argparse
import asyncio
import gzip
import hashlib
import logging
import os
import re
import tempfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Optional, Set, Tuple
from urllib.parse import urljoin, urlparse, urldefrag
from urllib.robotparser import RobotFileParser

import aiohttp
import chromadb
import PyPDF2
from bs4 import BeautifulSoup
from docx import Document

from rag_embeddings import Embedder, sanitize_metadata, fetch_all

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(levelname)s - %(message)s",
    handlers=[logging.FileHandler("crawler.log"), logging.StreamHandler()],
)
logger = logging.getLogger(__name__)

USER_AGENT = "CityWebsiteCrawler/1.0 (City Website RAG System)"


class MetadataExtractor:
    """Helper class for extracting structured metadata from documents.

    Document-type and department CLASSIFICATION was removed from the crawler:
    it is now done by classify.py (an LLM pass) for accuracy. What remains here
    is cheap factual extraction (dates, a PII heuristic, URL stats).
    """

    # Basic PII signal detection. Heuristic screen only, not authoritative.
    PII_PATTERNS = {
        "email": r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        "phone": r"\b\d{3}[-.]?\d{3}[-.]?\d{4}\b",
        "ssn": r"\b\d{3}-\d{2}-\d{4}\b",
        "address": r"\b\d+\s+[A-Za-z\s]+(?:Street|St|Avenue|Ave|Road|Rd|Boulevard|Blvd|Lane|Ln|Drive|Dr|Court|Ct)\b",
    }

    @staticmethod
    def extract_dates_from_text(text: str) -> Dict[str, Optional[str]]:
        dates = {"published_date_raw": None}
        date_patterns = [
            r"(?:published|posted|created|date):\s*(\d{1,2}[-/]\d{1,2}[-/]\d{2,4})",
            r"(?:published|posted|created|date):\s*([A-Za-z]+\s+\d{1,2},?\s+\d{4})",
        ]
        search_text = text[:1000]
        for pattern in date_patterns:
            matches = re.findall(pattern, search_text, re.IGNORECASE)
            if matches:
                dates["published_date_raw"] = matches[0]
                break
        return dates

    @staticmethod
    def extract_year_from_text(text: str) -> Optional[int]:
        year_matches = re.findall(r"\b(19\d{2}|20\d{2})\b", text[:500])
        if year_matches:
            return int(year_matches[0])
        return None

    @staticmethod
    def extract_fiscal_year(text: str) -> Optional[str]:
        fy_pattern = r"(?:FY|fiscal\s+year)\s*(\d{2,4})"
        matches = re.findall(fy_pattern, text[:500], re.IGNORECASE)
        if matches:
            year = matches[0]
            if len(year) == 2:
                year = f"20{year}" if int(year) < 50 else f"19{year}"
            return f"FY{year}"
        return None

    @staticmethod
    def detect_pii(text: str) -> Tuple[bool, List[str]]:
        detected_types = []
        sample_text = text[:5000]
        for pii_type, pattern in MetadataExtractor.PII_PATTERNS.items():
            if re.search(pattern, sample_text):
                detected_types.append(pii_type)
        return len(detected_types) > 0, detected_types

    @staticmethod
    def analyze_url(url: str) -> Dict:
        parsed = urlparse(url)
        path = parsed.path.strip("/")
        segments = [s for s in path.split("/") if s]
        return {
            "url_path": path,
            "url_depth": len(segments),
            "url_segments": segments,
            "file_extension": Path(parsed.path).suffix.lower().lstrip("."),
        }


class WebsiteCrawler:
    def __init__(
        self,
        base_url: str,
        db_path: str = "./chroma_db",
        collection_name: str = "city_website_content",
        embedding_model: str = "BAAI/bge-large-en-v1.5",
        max_file_size_mb: int = 100,
        max_concurrent: int = 10,
        chunk_size: int = 1000,
        chunk_overlap: int = 200,
        recreate: bool = False,
        max_depth: Optional[int] = None,
        max_pages: Optional[int] = None,
        allow_subdomains: bool = False,
        skip_pdf: bool = False,
        skip_spreadsheets: bool = False,
        exclude_patterns: Optional[List[str]] = None,
        include_patterns: Optional[List[str]] = None,
        sitemap: Optional[str] = None,
        device: str = "cpu",
    ):
        self.base_url = urldefrag(base_url).url.rstrip("/")
        self.domain = urlparse(base_url).netloc
        self.apex = self.domain[4:] if self.domain.startswith("www.") else self.domain
        self.max_depth = max_depth
        self.max_pages = max_pages
        self.allow_subdomains = allow_subdomains
        self.skip_pdf = skip_pdf
        self.skip_spreadsheets = skip_spreadsheets
        self.exclude_re = [re.compile(p) for p in (exclude_patterns or [])]
        self.include_re = [re.compile(p) for p in (include_patterns or [])]
        self.sitemap = sitemap
        self.db_path = db_path
        self.collection_name = collection_name
        self.max_file_size_bytes = max_file_size_mb * 1024 * 1024
        self.max_concurrent = max_concurrent
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.recreate = recreate

        self.crawled_urls: Set[str] = set()
        self.failed_urls: Set[str] = set()
        self.robots: Optional[RobotFileParser] = None

        self.metadata_extractor = MetadataExtractor()

        logger.info(f"Loading embedding model: {embedding_model} (first run downloads ~1.3GB)")
        self.embedder = Embedder(model_name=embedding_model, device=device)
        logger.info(f"Embedding dimension: {self.embedder.dimension}")

        # Local Chroma store. Cosine space matches normalized BGE vectors.
        self.chroma_client = chromadb.PersistentClient(path=db_path)
        if recreate:
            try:
                self.chroma_client.delete_collection(collection_name)
                logger.info(f"Deleted existing collection: {collection_name}")
            except Exception:
                pass
        self.collection = self.chroma_client.get_or_create_collection(
            name=collection_name, metadata={"hnsw:space": "cosine"}
        )

        self.session_http: Optional[aiohttp.ClientSession] = None

    async def __aenter__(self):
        timeout = aiohttp.ClientTimeout(total=30)
        connector = aiohttp.TCPConnector(limit=self.max_concurrent)
        self.session_http = aiohttp.ClientSession(
            timeout=timeout, connector=connector, headers={"User-Agent": USER_AGENT}
        )
        return self

    async def __aexit__(self, exc_type, exc_val, exc_tb):
        if self.session_http:
            await self.session_http.close()

    def chunk_text(self, text: str, title: str = "") -> List[str]:
        if len(text) <= self.chunk_size:
            return [text]

        chunks = []
        start = 0
        while start < len(text):
            end = start + self.chunk_size
            if end < len(text):
                last_period = text.rfind(".", start, end)
                last_question = text.rfind("?", start, end)
                last_exclamation = text.rfind("!", start, end)
                sentence_end = max(last_period, last_question, last_exclamation)
                if sentence_end > start + self.chunk_size // 2:
                    end = sentence_end + 1
                else:
                    last_space = text.rfind(" ", start, end)
                    if last_space > start + self.chunk_size // 2:
                        end = last_space

            chunk = text[start:end].strip()
            if chunk:
                if title and not chunk.startswith(title):
                    chunk = f"Title: {title}\n\n{chunk}"
                chunks.append(chunk)

            new_start = end - self.chunk_overlap
            start = new_start if new_start > start else end
        return chunks

    async def fetch_sitemap_urls(self, sitemap_url: str, _depth: int = 0,
                                 _seen: Optional[Set[str]] = None) -> Set[str]:
        """Fetch a sitemap (or sitemap index) and return the page URLs in it.

        Handles a <sitemapindex> by recursing into sub-sitemaps, and gzipped
        sub-sitemaps (.xml.gz). Recursion and count are bounded.
        """
        if _seen is None:
            _seen = set()
        if sitemap_url in _seen or _depth > 3 or len(_seen) > 200:
            return set()
        _seen.add(sitemap_url)

        urls: Set[str] = set()
        try:
            async with self.session_http.get(sitemap_url) as resp:
                if resp.status != 200:
                    logger.warning(f"Sitemap fetch failed {sitemap_url}: HTTP {resp.status}")
                    return urls
                raw = await resp.read()
            if raw[:2] == b"\x1f\x8b":  # gzip magic
                raw = gzip.decompress(raw)
            text = raw.decode("utf-8", "ignore")
            locs = re.findall(r"<loc>\s*(.*?)\s*</loc>", text, re.IGNORECASE | re.DOTALL)
            if "<sitemapindex" in text.lower():
                for sub in locs:
                    urls |= await self.fetch_sitemap_urls(sub.strip(), _depth + 1, _seen)
            else:
                for u in locs:
                    urls.add(urldefrag(u.strip()).url)
        except Exception as e:
            logger.warning(f"Error parsing sitemap {sitemap_url}: {e}")
        return urls

    async def check_robots_txt(self) -> RobotFileParser:
        robots_url = urljoin(self.base_url, "/robots.txt")
        rp = RobotFileParser()
        rp.set_url(robots_url)
        try:
            async with self.session_http.get(robots_url) as response:
                if response.status == 200:
                    robots_content = await response.text()
                    rp.parse(robots_content.splitlines())
                else:
                    rp.allow_all = True
        except Exception as e:
            logger.warning(f"Could not fetch robots.txt: {e}")
            rp.allow_all = True
        self.robots = rp
        return rp

    def is_valid_url(self, url: str) -> bool:
        # Drop #fragment: it is a within-page anchor, so page#a and page are the
        # same document. Defragging here keeps dedup and the frontier consistent.
        url = urldefrag(url).url
        parsed = urlparse(url)
        if self.allow_subdomains:
            if not (parsed.netloc == self.domain
                    or parsed.netloc == self.apex
                    or parsed.netloc.endswith("." + self.apex)):
                return False
        elif parsed.netloc != self.domain:
            return False

        skip_extensions = {".jpg", ".jpeg", ".png", ".gif", ".css", ".js", ".ico",
                           ".svg", ".woff", ".woff2", ".ttf", ".eot", ".zip", ".tar", ".gz"}
        path_lower = parsed.path.lower()
        if any(path_lower.endswith(ext) for ext in skip_extensions):
            return False

        if self.skip_pdf and path_lower.endswith(".pdf"):
            return False

        if self.skip_spreadsheets and path_lower.endswith(
                (".xls", ".xlsx", ".xlsm", ".csv", ".tsv", ".ods")):
            return False

        # URL pattern filters (regex against the full URL).
        if self.exclude_re and any(r.search(url) for r in self.exclude_re):
            return False
        if self.include_re and not any(r.search(url) for r in self.include_re):
            return False

        if url in self.crawled_urls:
            return False

        if self.robots is not None and not self.robots.can_fetch(USER_AGENT, url):
            logger.info(f"Disallowed by robots.txt: {url}")
            return False

        return True

    async def get_file_size(self, url: str) -> int:
        try:
            async with self.session_http.head(url) as response:
                content_length = response.headers.get("content-length")
                if content_length:
                    return int(content_length)
        except Exception:
            pass
        return 0

    async def extract_text_from_html(self, content: str, url: str) -> Tuple[str, str, List[str], Dict]:
        soup = BeautifulSoup(content, "html.parser")
        for script in soup(["script", "style", "nav", "footer", "header"]):
            script.decompose()

        title_tag = soup.find("title")
        title = title_tag.get_text().strip() if title_tag else urlparse(url).path

        main_content = (soup.find("main") or soup.find("article")
                        or soup.find("div", class_="content") or soup.body or soup)
        text = main_content.get_text(separator=" ", strip=True)
        text = re.sub(r"\s+", " ", text).strip()

        links = []
        for link in soup.find_all("a", href=True):
            absolute_url = urldefrag(urljoin(url, link["href"])).url
            if self.is_valid_url(absolute_url):
                links.append(absolute_url)

        metadata = {
            "has_tables": len(soup.find_all("table")) > 0,
            "has_images": len(soup.find_all("img")) > 0,
        }
        for meta_tag in soup.find_all("meta"):
            if meta_tag.get("property") == "article:published_time":
                metadata["published_date_raw"] = meta_tag.get("content")
            elif meta_tag.get("property") == "article:modified_time":
                metadata["modified_date_raw"] = meta_tag.get("content")
            elif meta_tag.get("name") == "date":
                metadata["published_date_raw"] = meta_tag.get("content")

        return text, title, links, metadata

    async def extract_text_from_pdf(self, content: bytes) -> Tuple[str, str, Dict]:
        metadata: Dict = {}
        try:
            with tempfile.NamedTemporaryFile() as tmp_file:
                tmp_file.write(content)
                tmp_file.flush()
                with open(tmp_file.name, "rb") as pdf_file:
                    pdf_reader = PyPDF2.PdfReader(pdf_file)
                    text = ""
                    title = ""
                    if pdf_reader.metadata:
                        title = pdf_reader.metadata.title or ""
                        metadata["author"] = pdf_reader.metadata.author
                        metadata["creator"] = pdf_reader.metadata.creator
                        metadata["producer"] = pdf_reader.metadata.producer
                        metadata["subject"] = pdf_reader.metadata.subject
                        if getattr(pdf_reader.metadata, "creation_date", None):
                            metadata["created_date_raw"] = pdf_reader.metadata.creation_date.isoformat()
                        if getattr(pdf_reader.metadata, "modification_date", None):
                            metadata["modified_date_raw"] = pdf_reader.metadata.modification_date.isoformat()
                    metadata["page_count"] = len(pdf_reader.pages)
                    for page in pdf_reader.pages:
                        text += (page.extract_text() or "") + "\n"
                    return text.strip(), title, metadata
        except Exception as e:
            logger.error(f"Error extracting PDF text: {e}")
            return "", "", metadata

    async def extract_text_from_docx(self, content: bytes) -> Tuple[str, str, Dict]:
        metadata: Dict = {}
        try:
            with tempfile.NamedTemporaryFile() as tmp_file:
                tmp_file.write(content)
                tmp_file.flush()
                doc = Document(tmp_file.name)
                title = ""
                if doc.core_properties.title:
                    title = doc.core_properties.title
                elif doc.paragraphs:
                    title = doc.paragraphs[0].text[:100]

                metadata["author"] = doc.core_properties.author
                metadata["creator"] = doc.core_properties.author
                metadata["subject"] = doc.core_properties.subject
                metadata["keywords"] = (
                    doc.core_properties.keywords.split(",") if doc.core_properties.keywords else []
                )
                if doc.core_properties.created:
                    metadata["created_date_raw"] = doc.core_properties.created.isoformat()
                if doc.core_properties.modified:
                    metadata["modified_date_raw"] = doc.core_properties.modified.isoformat()

                text = ""
                for paragraph in doc.paragraphs:
                    text += paragraph.text + "\n"
                metadata["has_tables"] = len(doc.tables) > 0
                return text.strip(), title, metadata
        except Exception as e:
            logger.error(f"Error extracting DOCX text: {e}")
            return "", "", metadata

    async def process_url(self, url: str, semaphore: asyncio.Semaphore) -> List[str]:
        async with semaphore:
            if url in self.crawled_urls or url in self.failed_urls:
                return []
            logger.info(f"Crawling: {url}")
            try:
                file_size = await self.get_file_size(url)
                if file_size > self.max_file_size_bytes:
                    logger.warning(f"Skipping {url} - file too large ({file_size} bytes)")
                    self.failed_urls.add(url)
                    return []

                async with self.session_http.get(url) as response:
                    if response.status != 200:
                        logger.warning(f"Failed to fetch {url}: HTTP {response.status}")
                        self.failed_urls.add(url)
                        return []

                    content_type = response.headers.get("content-type", "").lower()
                    last_modified = response.headers.get("last-modified")

                    if self.skip_spreadsheets and any(
                            t in content_type for t in
                            ("excel", "spreadsheetml", "ms-excel", "csv", "opendocument.spreadsheet")):
                        logger.info(f"Skipping spreadsheet (--skip-spreadsheets): {url}")
                        self.failed_urls.add(url)
                        return []

                    doc_metadata = {
                        "content_type_header": content_type,
                        "file_size": file_size if file_size > 0 else 0,
                    }
                    if last_modified:
                        doc_metadata["last_modified_header"] = last_modified

                    if "text/html" in content_type:
                        content = await response.text()
                        text, title, links, html_meta = await self.extract_text_from_html(content, url)
                        doc_metadata.update(html_meta)
                        file_type = "html"
                    elif "application/pdf" in content_type:
                        if self.skip_pdf:
                            logger.info(f"Skipping PDF (--skip-pdf): {url}")
                            self.failed_urls.add(url)
                            return []
                        content = await response.read()
                        text, title, pdf_meta = await self.extract_text_from_pdf(content)
                        doc_metadata.update(pdf_meta)
                        links = []
                        file_type = "pdf"
                    elif "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in content_type:
                        content = await response.read()
                        text, title, docx_meta = await self.extract_text_from_docx(content)
                        doc_metadata.update(docx_meta)
                        links = []
                        file_type = "docx"
                    else:
                        logger.warning(f"Unsupported content type for {url}: {content_type}")
                        self.failed_urls.add(url)
                        return []

                if not text or len(text.strip()) < 50:
                    logger.warning(f"No meaningful content extracted from {url}")
                    self.failed_urls.add(url)
                    return []

                if "published_date_raw" not in doc_metadata:
                    doc_metadata.update(self.metadata_extractor.extract_dates_from_text(text))

                doc_metadata["inferred_year"] = self.metadata_extractor.extract_year_from_text(f"{title} {text}")
                doc_metadata["fiscal_year"] = self.metadata_extractor.extract_fiscal_year(f"{title} {text}")

                has_pii, pii_types = self.metadata_extractor.detect_pii(text)
                doc_metadata["contains_potential_pii"] = has_pii
                doc_metadata["pii_types"] = pii_types

                doc_metadata.update(self.metadata_extractor.analyze_url(url))
                doc_metadata["content_length"] = len(text)
                doc_metadata["word_count"] = len(text.split())

                chunks = self.chunk_text(text, title)
                await self.store_chunks(url, title, chunks, file_type, doc_metadata)

                self.crawled_urls.add(url)
                logger.info(f"Processed {url} - {len(chunks)} chunks ({file_type})")
                return links if file_type == "html" else []
            except Exception as e:
                logger.error(f"Error processing {url}: {e}")
                self.failed_urls.add(url)
                return []

    async def store_chunks(self, url: str, title: str, chunks: List[str],
                           file_type: str, doc_metadata: Dict):
        """Embed chunks (off the event loop) and upsert them into Chroma."""
        crawled_at = datetime.now(timezone.utc).isoformat()

        # Embed all chunks of this page in one batched call, in a worker thread
        # so the CPU-bound encode does not stall other downloads.
        embeddings = await asyncio.to_thread(self.embedder.embed_documents, chunks)

        ids: List[str] = []
        documents: List[str] = []
        metadatas: List[Dict] = []

        for i, chunk in enumerate(chunks):
            doc_id = hashlib.md5(f"{url}_{i}".encode()).hexdigest()
            metadata = {
                "title": title,
                "url": url,
                "file_type": file_type,
                "chunk_id": i,
                "total_chunks": len(chunks),
                "content_hash": hashlib.md5(chunk.encode()).hexdigest(),
                "crawled_at": crawled_at,
            }
            metadata.update(doc_metadata)
            ids.append(doc_id)
            documents.append(chunk)
            metadatas.append(sanitize_metadata(metadata))

        # upsert: re-crawling the same URL overwrites prior chunks for it.
        await asyncio.to_thread(
            self.collection.upsert,
            ids=ids, embeddings=embeddings, documents=documents, metadatas=metadatas,
        )

    async def crawl_website(self):
        logger.info(f"Starting crawl of {self.base_url}")
        if self.max_depth is not None:
            logger.info(f"Max depth: {self.max_depth}")
        if self.max_pages is not None:
            logger.info(f"Max pages: {self.max_pages}")
        await self.check_robots_txt()

        # Frontier maps url -> depth. Seeds are depth 0.
        if self.sitemap:
            found = await self.fetch_sitemap_urls(self.sitemap)
            seeds = [u for u in found if self.is_valid_url(u)]
            logger.info(f"Sitemap: {len(found)} URLs listed, {len(seeds)} kept after filters")
            if not seeds:
                logger.warning("No usable sitemap URLs after filtering; nothing to crawl")
                self.print_crawl_summary()
                return
            frontier: Dict[str, int] = {u: 0 for u in seeds}
        else:
            frontier = {self.base_url: 0}
        in_flight: Set[str] = set()
        semaphore = asyncio.Semaphore(self.max_concurrent)

        while True:
            candidates = [(u, d) for u, d in frontier.items()
                          if u not in self.crawled_urls and u not in self.failed_urls
                          and u not in in_flight]
            if not candidates:
                break

            # Stop at max_pages, and never start more than the remaining budget.
            if self.max_pages is not None:
                remaining = self.max_pages - len(self.crawled_urls)
                if remaining <= 0:
                    break
                candidates = candidates[:remaining]

            batch_urls = [u for u, _ in candidates]
            depth_of = {u: d for u, d in candidates}
            in_flight.update(batch_urls)
            results = await asyncio.gather(
                *[self.process_url(u, semaphore) for u in batch_urls],
                return_exceptions=True,
            )
            in_flight.difference_update(batch_urls)
            for u in batch_urls:
                frontier.pop(u, None)

            for u, result in zip(batch_urls, results):
                if isinstance(result, list):
                    next_depth = depth_of[u] + 1
                    # Do not enqueue links deeper than max_depth.
                    if self.max_depth is not None and next_depth > self.max_depth:
                        continue
                    for new_url in result:
                        if new_url in self.crawled_urls or new_url in self.failed_urls:
                            continue
                        if new_url not in frontier or next_depth < frontier[new_url]:
                            frontier[new_url] = next_depth

            logger.info(
                f"Batch complete. Crawled: {len(self.crawled_urls)}, "
                f"Failed: {len(self.failed_urls)}, Frontier: {len(frontier)}"
            )
            await asyncio.sleep(1)

        logger.info(
            f"Crawl complete. Crawled {len(self.crawled_urls)} URLs, failed on {len(self.failed_urls)}"
        )
        self.print_crawl_summary()

    def print_crawl_summary(self):
        """Aggregate Chroma metadata in Python (no server-side aggregations)."""
        logger.info("=" * 70)
        logger.info("CRAWL SUMMARY")
        logger.info("(document type / department are filled in later by classify.py)")
        logger.info("=" * 70)
        try:
            metadatas = fetch_all(self.collection, ["metadatas"])["metadatas"]
            logger.info(f"Total chunks indexed: {len(metadatas)}")

            by_file = Counter(m.get("file_type", "unknown") for m in metadatas)
            by_year = Counter(m.get("inferred_year") for m in metadatas if m.get("inferred_year"))
            pii_count = sum(1 for m in metadatas if m.get("contains_potential_pii"))

            logger.info("File Types:")
            for key, count in by_file.most_common():
                logger.info(f"  {str(key):10s}: {count}")
            logger.info("Years Found:")
            for key in sorted(by_year):
                logger.info(f"  {key}: {by_year[key]}")
            logger.info(f"Chunks with potential PII: {pii_count}")
        except Exception as e:
            logger.error(f"Error generating summary: {e}")
        logger.info("=" * 70)


def _env_default(name, fallback):
    return os.environ.get(name, fallback)


async def main():
    parser = argparse.ArgumentParser(description="Local crawler (Chroma + BGE-large)")
    parser.add_argument("--base-url", required=True, help="Base URL to crawl")
    parser.add_argument("--db-path", default=_env_default("CHROMA_PATH", "./chroma_db"),
                        help="Chroma persistence directory (must match the chatbot)")
    parser.add_argument("--collection-name", default=_env_default("CHROMA_COLLECTION", "city_website_content"))
    parser.add_argument("--embedding-model", default=_env_default("EMBEDDING_MODEL", "BAAI/bge-large-en-v1.5"))
    parser.add_argument("--max-file-size-mb", type=int, default=100)
    parser.add_argument("--max-concurrent", type=int, default=10)
    parser.add_argument("--chunk-size", type=int, default=1000)
    parser.add_argument("--chunk-overlap", type=int, default=200)
    parser.add_argument("--recreate", action="store_true",
                        help="Delete and rebuild the collection instead of upserting")
    parser.add_argument("--max-depth", type=int, default=None,
                        help="Max link depth from the start URL (0 = only the start page). Unlimited if unset.")
    parser.add_argument("--max-pages", type=int, default=None,
                        help="Stop after crawling this many pages. Unlimited if unset.")
    parser.add_argument("--allow-subdomains", action="store_true",
                        help="Also crawl subdomains of the start URL's domain (default: same host only)")
    parser.add_argument("--skip-pdf", action="store_true",
                        help="Do not fetch or process PDF files")
    parser.add_argument("--skip-spreadsheets", action="store_true",
                        help="Do not fetch or process spreadsheets (xls, xlsx, xlsm, csv, tsv, ods)")
    parser.add_argument("--exclude-pattern", action="append", default=None, metavar="REGEX",
                        help="Skip URLs matching this regex (repeatable). "
                             "Example: --exclude-pattern 'annual_permits' --exclude-pattern '\\?year='")
    parser.add_argument("--include-pattern", action="append", default=None, metavar="REGEX",
                        help="Only crawl URLs matching at least one of these regexes (repeatable)")
    parser.add_argument("--sitemap", default=None, metavar="URL",
                        help="Seed the crawl from this sitemap URL (handles sitemap index and .gz). "
                             "Combine with --max-depth 1 to also branch one hop off listed pages.")
    parser.add_argument("--device", default="cpu",
                        help="Embedding device: 'cpu' or 'cuda' (cuda requires a GPU build of torch)")
    args = parser.parse_args()

    async with WebsiteCrawler(
        base_url=args.base_url,
        db_path=args.db_path,
        collection_name=args.collection_name,
        embedding_model=args.embedding_model,
        max_file_size_mb=args.max_file_size_mb,
        max_concurrent=args.max_concurrent,
        chunk_size=args.chunk_size,
        chunk_overlap=args.chunk_overlap,
        recreate=args.recreate,
        max_depth=args.max_depth,
        max_pages=args.max_pages,
        allow_subdomains=args.allow_subdomains,
        skip_pdf=args.skip_pdf,
        skip_spreadsheets=args.skip_spreadsheets,
        exclude_patterns=args.exclude_pattern,
        include_patterns=args.include_pattern,
        sitemap=args.sitemap,
        device=args.device,
    ) as crawler:
        await crawler.crawl_website()


if __name__ == "__main__":
    asyncio.run(main())
