#!/usr/bin/env python3
"""
Shared text extraction and chunking for the Gigi City RAG system.

Single source of truth for turning raw bytes/markup into the normalized text and
chunks that get embedded. The crawler calls these at index time; the V3 backend
calls the SAME functions at verify time. If extraction or chunk boundaries drift
between the two, verification quotes stop lining up with stored content, so this
logic must live in exactly one place (see V3.md sections 2.5 and 9).

Normalization contract (must not change without re-indexing):
  * HTML: strip script/style/nav/footer/header, take main/article/.content/body,
    join across elements with single spaces, collapse whitespace with
    re.sub(r"\\s+", " ", text).strip().
  * chunk_text injects a ``Title: ...`` prefix into MULTI-chunk documents only.
    That prefix is NOT present on the live page and must never be used as a
    highlight quote (relevant to V3 verification/snap and the companion).

Heavy parsers (bs4, PyPDF2, python-docx) are imported lazily inside each function
so importing this module stays cheap for callers that only need chunk_text.
"""

import re
import tempfile
from typing import Dict, List, Tuple
from urllib.parse import urlparse


def extract_text_from_html(content: str, url: str = "") -> Tuple[str, str, Dict]:
    """Extract (normalized_text, title, metadata) from HTML.

    Link discovery is intentionally NOT done here: it needs crawler-specific
    URL-validity state, so the crawler parses links separately (over a soup
    decomposed the same way) while reusing this for text/title/metadata.
    """
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(content, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()

    title_tag = soup.find("title")
    title = title_tag.get_text().strip() if title_tag else (urlparse(url).path if url else "")

    main_content = (soup.find("main") or soup.find("article")
                    or soup.find("div", class_="content") or soup.body or soup)
    text = main_content.get_text(separator=" ", strip=True)
    text = re.sub(r"\s+", " ", text).strip()

    metadata: Dict = {
        "has_tables": len(soup.find_all("table")) > 0,
        "has_images": len(soup.find_all("img")) > 0,
    }
    for meta_tag in soup.find_all("meta"):
        if meta_tag.get("property") == "article:published_time":
            metadata["published_date_raw"] = meta_tag.get("content")
        elif meta_tag.get("property") == "article:modified_time":
            metadata["modified_date_raw"] = meta_tag.get("content")
        elif meta_tag.get("name") == "date":
            metadata["published_date_raw"] = meta_tag.get("content")

    return text, title, metadata


def extract_text_from_pdf(content: bytes) -> Tuple[str, str, Dict]:
    """Extract (text, title, metadata) from PDF bytes. Pages joined with newlines."""
    import PyPDF2

    metadata: Dict = {}
    try:
        with tempfile.NamedTemporaryFile() as tmp:
            tmp.write(content)
            tmp.flush()
            with open(tmp.name, "rb") as fh:
                reader = PyPDF2.PdfReader(fh)
                text = ""
                title = ""
                if reader.metadata:
                    title = reader.metadata.title or ""
                    metadata["author"] = reader.metadata.author
                    metadata["creator"] = reader.metadata.creator
                    metadata["producer"] = reader.metadata.producer
                    metadata["subject"] = reader.metadata.subject
                    if getattr(reader.metadata, "creation_date", None):
                        metadata["created_date_raw"] = reader.metadata.creation_date.isoformat()
                    if getattr(reader.metadata, "modification_date", None):
                        metadata["modified_date_raw"] = reader.metadata.modification_date.isoformat()
                metadata["page_count"] = len(reader.pages)
                for page in reader.pages:
                    text += (page.extract_text() or "") + "\n"
                return text.strip(), title, metadata
    except Exception as e:
        import logging
        logging.getLogger(__name__).error(f"Error extracting PDF text: {e}")
        return "", "", metadata


def extract_text_from_docx(content: bytes) -> Tuple[str, str, Dict]:
    """Extract (text, title, metadata) from DOCX bytes. Paragraphs joined with newlines."""
    from docx import Document

    metadata: Dict = {}
    try:
        with tempfile.NamedTemporaryFile(suffix=".docx") as tmp:
            tmp.write(content)
            tmp.flush()
            doc = Document(tmp.name)
            title = ""
            if doc.core_properties.title:
                title = doc.core_properties.title
            elif doc.paragraphs:
                title = doc.paragraphs[0].text[:100]

            metadata["author"] = doc.core_properties.author
            metadata["creator"] = doc.core_properties.author
            metadata["subject"] = doc.core_properties.subject
            metadata["keywords"] = (
                doc.core_properties.keywords.split(",") if doc.core_properties.keywords else []
            )
            if doc.core_properties.created:
                metadata["created_date_raw"] = doc.core_properties.created.isoformat()
            if doc.core_properties.modified:
                metadata["modified_date_raw"] = doc.core_properties.modified.isoformat()

            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            metadata["has_tables"] = len(doc.tables) > 0
            return text.strip(), title, metadata
    except Exception as e:
        import logging
        logging.getLogger(__name__).error(f"Error extracting DOCX text: {e}")
        return "", "", metadata


def extract_text_from_plain(content: bytes) -> Tuple[str, str, Dict]:
    """TXT / MD. No real metadata; title is left to the caller / first line."""
    text = content.decode("utf-8", "ignore").strip()
    return text, "", {}


def chunk_text(text: str, title: str = "", chunk_size: int = 1000,
               chunk_overlap: int = 200) -> List[str]:
    """Split text into overlapping, sentence-aware chunks.

    Behavior is the crawler's canonical logic: short documents (<= chunk_size)
    return a single unprefixed chunk; longer documents are split on sentence
    boundaries (then whitespace) and each chunk gets a ``Title:`` prefix. Keeping
    one implementation guarantees the crawler, add_resource, and the backend
    produce identical chunk boundaries.
    """
    if len(text) <= chunk_size:
        return [text]

    chunks: List[str] = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        if end < len(text):
            last_period = text.rfind(".", start, end)
            last_question = text.rfind("?", start, end)
            last_exclamation = text.rfind("!", start, end)
            sentence_end = max(last_period, last_question, last_exclamation)
            if sentence_end > start + chunk_size // 2:
                end = sentence_end + 1
            else:
                last_space = text.rfind(" ", start, end)
                if last_space > start + chunk_size // 2:
                    end = last_space

        chunk = text[start:end].strip()
        if chunk:
            if title and not chunk.startswith(title):
                chunk = f"Title: {title}\n\n{chunk}"
            chunks.append(chunk)

        new_start = end - chunk_overlap
        start = new_start if new_start > start else end
    return chunks
