#!/usr/bin/env python3
"""
Add a single resource (local file or one URL) to the existing Chroma collection.

Companion to website_crawler.py. The crawler walks a whole site; this adds ONE
thing at a time, for knowledge augmentation:

  * A local file you want Gigi to know about (PDF, DOCX, HTML, TXT, MD).
    Optionally give it a --url so Gigi can link residents to it.
  * A single URL to fetch and ingest without running a whole crawl.

It writes chunks in the SAME shape as the crawler (same metadata keys, same
chunking, same ID scheme when a URL is present) so retrieval and the chatbot's
citation/footer logic treat ingested resources identically to crawled pages.

IDs:
  * With --url   : md5(f"{url}_{i}")        -- matches the crawler, so this
                   upserts over a previously crawled copy of the same URL.
  * Without --url: md5(f"local:{source_id}_{i}"), where source_id defaults to
                   the absolute file path (override with --source-id). Re-running
                   the same file upserts instead of duplicating.

PII: the same heuristic the crawler uses runs here. If it fires, ingest is
REFUSED unless you pass --allow-pii. Anything ingested can be sent to the
Anthropic API at chat (and classify) time, so resident PII handling falls under
City of Garden Grove Administrative Regulation 2.16 (Cloud Computing Services
Policy): https://internal.ggcity.org/policies

Set your key only if you also classify later; ingestion itself needs no API key.

Examples:
  # A public PDF you want Gigi to cite and link:
  python add_resource.py --file ./fee_schedule_2026.pdf \
      --url https://www.ggcity.org/finance/fee-schedule-2026.pdf

  # A local-only knowledge file with no public URL:
  python add_resource.py --file ./internal_faq.md --title "Water Billing FAQ"

  # A single web page, no crawl:
  python add_resource.py --url https://www.ggcity.org/trash/bulk-pickup
"""

import argparse
import hashlib
import logging
import os
import re
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Optional, Tuple
from urllib.parse import urldefrag, urlparse

import chromadb

from rag_embeddings import Embedder, sanitize_metadata

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)

USER_AGENT = "CityResourceIngest/1.0 (City Website RAG System)"

# Same PII heuristic the crawler uses. Kept in sync deliberately.
PII_PATTERNS = {
    "email": r"\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
    "phone": r"\b\d{3}[-.]?\d{3}[-.]?\d{4}\b",
    "ssn": r"\b\d{3}-\d{2}-\d{4}\b",
    "address": r"\b\d+\s+[A-Za-z\s]+(?:Street|St|Avenue|Ave|Road|Rd|Boulevard|Blvd|"
               r"Lane|Ln|Drive|Dr|Court|Ct)\b",
}


# --------------------------------------------------------------------------
# Extraction. Standalone copies of the crawler's logic (the crawler's are
# instance methods tangled with crawler state). Behavior matches the crawler.
# --------------------------------------------------------------------------
def extract_text_from_pdf(content: bytes) -> Tuple[str, str, Dict]:
    import PyPDF2
    metadata: Dict = {}
    try:
        with tempfile.NamedTemporaryFile() as tmp:
            tmp.write(content)
            tmp.flush()
            with open(tmp.name, "rb") as fh:
                reader = PyPDF2.PdfReader(fh)
                text = ""
                title = ""
                if reader.metadata:
                    title = reader.metadata.title or ""
                    metadata["author"] = reader.metadata.author
                    metadata["creator"] = reader.metadata.creator
                    metadata["producer"] = reader.metadata.producer
                    metadata["subject"] = reader.metadata.subject
                    if getattr(reader.metadata, "creation_date", None):
                        metadata["created_date_raw"] = reader.metadata.creation_date.isoformat()
                    if getattr(reader.metadata, "modification_date", None):
                        metadata["modified_date_raw"] = reader.metadata.modification_date.isoformat()
                metadata["page_count"] = len(reader.pages)
                for page in reader.pages:
                    text += (page.extract_text() or "") + "\n"
                return text.strip(), title, metadata
    except Exception as e:
        logger.error(f"Error extracting PDF text: {e}")
        return "", "", metadata


def extract_text_from_docx(content: bytes) -> Tuple[str, str, Dict]:
    from docx import Document
    metadata: Dict = {}
    try:
        with tempfile.NamedTemporaryFile(suffix=".docx") as tmp:
            tmp.write(content)
            tmp.flush()
            doc = Document(tmp.name)
            title = ""
            if doc.core_properties.title:
                title = doc.core_properties.title
            elif doc.paragraphs:
                title = doc.paragraphs[0].text[:100]
            metadata["author"] = doc.core_properties.author
            metadata["creator"] = doc.core_properties.author
            metadata["subject"] = doc.core_properties.subject
            metadata["keywords"] = (
                doc.core_properties.keywords.split(",") if doc.core_properties.keywords else []
            )
            if doc.core_properties.created:
                metadata["created_date_raw"] = doc.core_properties.created.isoformat()
            if doc.core_properties.modified:
                metadata["modified_date_raw"] = doc.core_properties.modified.isoformat()
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            metadata["has_tables"] = len(doc.tables) > 0
            return text.strip(), title, metadata
    except Exception as e:
        logger.error(f"Error extracting DOCX text: {e}")
        return "", "", metadata


def extract_text_from_html(content: str, url: str = "") -> Tuple[str, str, Dict]:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(content, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    title_tag = soup.find("title")
    title = title_tag.get_text().strip() if title_tag else (urlparse(url).path if url else "")
    main = (soup.find("main") or soup.find("article")
            or soup.find("div", class_="content") or soup.body or soup)
    text = main.get_text(separator=" ", strip=True)
    text = re.sub(r"\s+", " ", text).strip()
    metadata = {
        "has_tables": len(soup.find_all("table")) > 0,
        "has_images": len(soup.find_all("img")) > 0,
    }
    for meta_tag in soup.find_all("meta"):
        if meta_tag.get("property") == "article:published_time":
            metadata["published_date_raw"] = meta_tag.get("content")
        elif meta_tag.get("property") == "article:modified_time":
            metadata["modified_date_raw"] = meta_tag.get("content")
        elif meta_tag.get("name") == "date":
            metadata["published_date_raw"] = meta_tag.get("content")
    return text, title, metadata


def extract_text_from_plain(content: bytes) -> Tuple[str, str, Dict]:
    """TXT / MD. No real metadata; title is left to caller / first line."""
    text = content.decode("utf-8", "ignore").strip()
    return text, "", {}


# --------------------------------------------------------------------------
# Chunking. Copied VERBATIM from website_crawler.WebsiteCrawler.chunk_text so
# chunk boundaries match crawled content exactly.
# --------------------------------------------------------------------------
def chunk_text(text: str, title: str, chunk_size: int, chunk_overlap: int) -> List[str]:
    if len(text) <= chunk_size:
        single = text.strip()
        if single and title and not single.startswith(title):
            single = f"Title: {title}\n\n{single}"
        return [single] if single else []

    chunks: List[str] = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        if end < len(text):
            last_period = text.rfind(".", start, end)
            last_question = text.rfind("?", start, end)
            last_exclamation = text.rfind("!", start, end)
            sentence_end = max(last_period, last_question, last_exclamation)
            if sentence_end > start + chunk_size // 2:
                end = sentence_end + 1
            else:
                last_space = text.rfind(" ", start, end)
                if last_space > start + chunk_size // 2:
                    end = last_space
        chunk = text[start:end].strip()
        if chunk:
            if title and not chunk.startswith(title):
                chunk = f"Title: {title}\n\n{chunk}"
            chunks.append(chunk)
        new_start = end - chunk_overlap
        start = new_start if new_start > start else end
    return chunks


# --------------------------------------------------------------------------
# PII + metadata helpers (match the crawler).
# --------------------------------------------------------------------------
def detect_pii(text: str) -> Tuple[bool, List[str]]:
    detected = []
    sample = text[:5000]
    for pii_type, pattern in PII_PATTERNS.items():
        if re.search(pattern, sample):
            detected.append(pii_type)
    return len(detected) > 0, detected


def extract_year_from_text(text: str) -> Optional[int]:
    matches = re.findall(r"\b(19\d{2}|20\d{2})\b", text[:500])
    return int(matches[0]) if matches else None


def extract_fiscal_year(text: str) -> Optional[str]:
    matches = re.findall(r"(?:FY|fiscal\s+year)\s*(\d{2,4})", text[:500], re.IGNORECASE)
    if matches:
        year = matches[0]
        if len(year) == 2:
            year = f"20{year}" if int(year) < 50 else f"19{year}"
        return f"FY{year}"
    return None


def analyze_url(url: str) -> Dict:
    parsed = urlparse(url)
    path = parsed.path.strip("/")
    segments = [s for s in path.split("/") if s]
    return {
        "url_path": path,
        "url_depth": len(segments),
        "url_segments": segments,
        "file_extension": Path(parsed.path).suffix.lower().lstrip("."),
    }


# --------------------------------------------------------------------------
# Source acquisition.
# --------------------------------------------------------------------------
def guess_file_type(path_or_url: str, content_type: str = "") -> str:
    ct = content_type.lower()
    if "application/pdf" in ct:
        return "pdf"
    if "wordprocessingml.document" in ct:
        return "docx"
    if "text/html" in ct:
        return "html"
    if "text/plain" in ct or "markdown" in ct:
        return "txt"
    ext = Path(urlparse(path_or_url).path if "://" in path_or_url else path_or_url).suffix.lower()
    return {
        ".pdf": "pdf", ".docx": "docx", ".html": "html", ".htm": "html",
        ".txt": "txt", ".md": "txt", ".markdown": "txt",
    }.get(ext, "")


def load_local_file(path: str) -> Tuple[bytes, str, str]:
    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError(f"No such file: {path}")
    return p.read_bytes(), guess_file_type(path), str(p.resolve())


def fetch_url(url: str) -> Tuple[bytes, str, str]:
    import requests
    headers = {"User-Agent": USER_AGENT}
    resp = requests.get(url, headers=headers, timeout=30)
    resp.raise_for_status()
    ct = resp.headers.get("content-type", "")
    return resp.content, guess_file_type(url, ct), ct


def extract(content: bytes, file_type: str, url: str = "") -> Tuple[str, str, Dict]:
    if file_type == "pdf":
        return extract_text_from_pdf(content)
    if file_type == "docx":
        return extract_text_from_docx(content)
    if file_type == "html":
        return extract_text_from_html(content.decode("utf-8", "ignore"), url)
    if file_type == "txt":
        return extract_text_from_plain(content)
    raise ValueError(f"Unsupported file type: {file_type!r}")


# --------------------------------------------------------------------------
# Main ingest.
# --------------------------------------------------------------------------
def ingest(args) -> int:
    if not args.file and not args.url:
        logger.error("Give --file, --url, or both.")
        return 2

    # 1. Acquire bytes.
    try:
        if args.file:
            content, file_type, abs_path = load_local_file(args.file)
            origin = abs_path
        else:
            content, file_type, _ = fetch_url(args.url)
            origin = args.url
    except Exception as e:
        logger.error(f"Could not load source: {e}")
        return 1

    if args.type:
        file_type = args.type
    if not file_type:
        logger.error("Could not determine file type; pass --type "
                     "(pdf|docx|html|txt).")
        return 1

    # 2. Extract text + native metadata.
    try:
        text, extracted_title, doc_meta = extract(content, file_type, args.url or "")
    except Exception as e:
        logger.error(f"Extraction failed: {e}")
        return 1

    if not text or len(text.strip()) < 50:
        logger.error(f"No meaningful text extracted from {origin} "
                     f"({len(text.strip())} chars). Nothing ingested.")
        return 1

    title = args.title or extracted_title or (
        Path(args.file).stem if args.file else urlparse(args.url).path
    )

    # 3. PII gate.
    has_pii, pii_types = detect_pii(text)
    if has_pii and not args.allow_pii:
        logger.error(
            "PII heuristic flagged this source (%s). Refusing to ingest.\n"
            "If this content is cleared for the cloud-reachable store under "
            "AR 2.16 (https://internal.ggcity.org/policies), re-run with "
            "--allow-pii.", ", ".join(pii_types),
        )
        return 3
    if has_pii:
        logger.warning("PII heuristic flagged (%s); ingesting anyway because "
                       "--allow-pii was set.", ", ".join(pii_types))

    # 4. URL handling + ID basis.
    url = urldefrag(args.url).url if args.url else ""
    if url:
        id_basis = url
    else:
        source_id = args.source_id or origin  # absolute path by default
        id_basis = f"local:{source_id}"
        logger.warning("No --url given: this source will have no clickable "
                       "citation link. Gigi can use its content but cannot link "
                       "residents to it.")

    # 5. Chunk (verbatim crawler logic).
    chunks = chunk_text(text, title, args.chunk_size, args.chunk_overlap)
    if not chunks:
        logger.error("Chunking produced nothing. Nothing ingested.")
        return 1

    # 6. Build metadata matching the crawler's shape.
    base_meta: Dict = dict(doc_meta)
    base_meta["inferred_year"] = extract_year_from_text(f"{title} {text}")
    base_meta["fiscal_year"] = extract_fiscal_year(f"{title} {text}")
    base_meta["contains_potential_pii"] = has_pii
    base_meta["pii_types"] = pii_types
    base_meta["content_length"] = len(text)
    base_meta["word_count"] = len(text.split())
    if url:
        base_meta.update(analyze_url(url))
    base_meta["ingested_via"] = "add_resource"  # provenance: distinguishes from crawl
    if args.file:
        base_meta["local_source_path"] = origin

    # 7. Embed + upsert.
    logger.info("Loading embedding model: %s", args.embedding_model)
    embedder = Embedder(model_name=args.embedding_model, device=args.device)
    embeddings = embedder.embed_documents(chunks)

    client = chromadb.PersistentClient(path=args.db_path)
    collection = client.get_or_create_collection(
        name=args.collection_name, metadata={"hnsw:space": "cosine"})

    crawled_at = datetime.now(timezone.utc).isoformat()
    ids, documents, metadatas = [], [], []
    for i, chunk in enumerate(chunks):
        doc_id = hashlib.md5(f"{id_basis}_{i}".encode()).hexdigest()
        meta = {
            "title": title,
            "url": url,
            "file_type": file_type,
            "chunk_id": i,
            "total_chunks": len(chunks),
            "content_hash": hashlib.md5(chunk.encode()).hexdigest(),
            "crawled_at": crawled_at,
        }
        meta.update(base_meta)
        ids.append(doc_id)
        documents.append(chunk)
        metadatas.append(sanitize_metadata(meta))

    if args.dry_run:
        logger.info("DRY RUN: would upsert %d chunks for %r (title=%r, url=%r, "
                    "file_type=%s).", len(chunks), id_basis, title, url, file_type)
        logger.info("First chunk preview:\n%s", chunks[0][:500])
        return 0

    collection.upsert(ids=ids, embeddings=embeddings,
                      documents=documents, metadatas=metadatas)
    logger.info("Ingested %d chunks for %r (title=%r, url=%r). Collection now has "
                "%d chunks.", len(chunks), id_basis, title, url or "(none)",
                collection.count())
    if not url:
        logger.info("Reminder: no URL, so this appears in retrieval but not as a "
                    "clickable source. classify.py will label it on its next run.")
    return 0


def _env_default(name, fallback):
    return os.environ.get(name, fallback)


def main():
    p = argparse.ArgumentParser(
        description="Add one local file or URL to the existing Chroma collection.")
    src = p.add_argument_group("source (give --file, --url, or both)")
    src.add_argument("--file", help="Path to a local file (pdf/docx/html/txt/md)")
    src.add_argument("--url", help="Citation URL. If --file is also given, the URL "
                                   "is used for linking/IDs but the file's bytes "
                                   "are ingested. If only --url, it is fetched.")
    src.add_argument("--type", choices=["pdf", "docx", "html", "txt"],
                     help="Force the file type instead of guessing from extension/header")
    src.add_argument("--title", help="Override the document title used in citations")
    src.add_argument("--source-id",
                     help="Stable ID basis for a no-URL local file (default: "
                          "absolute file path). Use to keep IDs stable if the file moves.")

    p.add_argument("--allow-pii", action="store_true",
                   help="Ingest even if the PII heuristic fires (default: refuse). "
                        "Subject to AR 2.16.")

    db = p.add_argument_group("store (must match the crawler/chatbot)")
    db.add_argument("--db-path", default=_env_default("CHROMA_PATH", "./chroma_db"))
    db.add_argument("--collection-name",
                    default=_env_default("CHROMA_COLLECTION", "city_website_content"))
    db.add_argument("--embedding-model",
                    default=_env_default("EMBEDDING_MODEL", "BAAI/bge-large-en-v1.5"),
                    help="MUST match the crawler, or retrieval breaks")
    db.add_argument("--device", default="cpu", help="'cpu' or 'cuda'")
    db.add_argument("--chunk-size", type=int, default=1000,
                    help="Match the crawler (default 1000) so chunks line up")
    db.add_argument("--chunk-overlap", type=int, default=200,
                    help="Match the crawler (default 200)")

    p.add_argument("--dry-run", action="store_true",
                   help="Extract, chunk, and report, but do not write to Chroma")
    args = p.parse_args()

    sys.exit(ingest(args))


if __name__ == "__main__":
    main()
